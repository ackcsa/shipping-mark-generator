#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
唛头文件生成器 (Shipping Mark Generator)

输入：一张模板 xlsx —— 明细行**与 PL箱单 对齐**，列名同 PL箱单，从箱单复制粘贴即可
    列（A..M，与 PL箱单 同名同序）：
        Items | Descriptions | Specifications | Material | Quantity | Unit | Packages |
        NW | Net weight (kg) | GW | Gross weight (kg) | CBM | 品名中英文
    其右追加（唛头专用，PL 里没有，需手填）：
        票名 | 厂家英文 | 发货人 | 发货人地址 | 收货人 | 收货人地址 | 原产国
    行：r1 = 表头；r2 起 = 明细行（从 PL箱单 明细区复制粘贴）

输出：一份合并 docx —— 一件一个唛头、一页一张、块间分页符。

编号规则：**按行从上到下累加**。每行的 Packages 就是该行要出的唛头个数；
         PACKAGE NO = <本票件数合计>-<序号>，序号跨行连续 1,2,3…；
         本票所有唛头的 TOTAL 都 = 本票件数合计。

做法：以母本 docx（示例唛头）为版式基准，取第 1 块 11 段作块模板，
      只替换各段「值 run」的文字，其余 XML（pPr / rPr / 字体 / 字号 / 加粗 / 页面设置）原样保留；
      其余 zip 部件原样搬运，只重写 word/document.xml 与 docProps/app.xml。

用法：
    python make_mark.py                          # 用同目录 mark_config.json
    python make_mark.py -c 别的配置.json
    python make_mark.py -i 模板表.xlsx -o 输出.docx
    python make_mark.py --no-pl-check            # 跳过件数 vs 三单/四单 PL 校验
    python make_mark.py --dry-run                # 只解析+校验并打印，不写文件
"""

import argparse
import glob
import json
import math
import os
import re
import subprocess
import sys
import zipfile

# 中文 Windows 的 cmd 默认 936 代码页 —— 控制台输出不能因编码问题崩掉
def _make_console_safe():
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(errors="replace")
        except Exception:
            pass


_make_console_safe()


APP_NAME = "唛头生成器"
APP_VERSION = "2.0.0"


def base_dir():
    """程序所在目录：打包成 exe 后 = exe 所在目录，否则 = 脚本所在目录。"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


def bundle_dir():
    """打包进 exe 的随附资源目录（onefile 下 = %TEMP%\\_MEIxxxx）；未打包 = 脚本目录。

    ★ 只用来做**兜底**：正常情况下配置/母本都在 exe 旁边，用户改了就地生效。
    只有在用户把 exe 单独拷走、旁边什么都没有时，才回落到打包进 exe 的那份。
    """
    meipass = getattr(sys, "_MEIPASS", None)
    return meipass if meipass else os.path.dirname(os.path.abspath(__file__))


def find_asset(name, prefer_dir=None):
    """按「用户旁边那份 → 打包进 exe 的那份」顺序找一个随附文件，都没有则返回 None。"""
    for d in (prefer_dir, base_dir(), bundle_dir()):
        if not d:
            continue
        p = os.path.join(d, name)
        if os.path.exists(p):
            return p
    return None

# ---------------------------------------------------------------- 常量

# 一个唛头的 11 行；(行序, 说明, 模板 run 数, 需要替换的 run 下标)
LINE_SPECS = [
    (0,  "SHIPPER",            2, [1]),
    (1,  "SHIPPER'S ADDRESS",  2, [1]),
    (2,  "(空行)",             0, []),
    (3,  "CONSIGNEE",          2, [1]),
    (4,  "CNEE'S ADDRESS",     2, [1]),
    (5,  "(空行)",             0, []),
    (6,  "COMMODITY",          2, [1]),
    (7,  "MANUFACTURER",       2, [1]),
    (8,  "ORIGIN",             2, [1]),
    (9,  "TOTAL",              2, [1]),
    (10, "PACKAGE NO",         4, [1, 3]),
]

BLOCK_LINES = 11

# 母本块模板各段应当出现的标签（用于校验母本没被换掉）
EXPECTED_LABELS = ["SHIPPER:", "SHIPPER", None, "CONSIGNEE:", "CNEE", None,
                   "COMMODITY:", "MANUFACTURER:", "ORIGIN:", "TOTAL:", "PACKAGE"]

# 唛头要用到的字段 -> 输入表表头可用的写法
FIELD_ALIASES = {
    "票名":       ["票名", "合同号", "票号", "合同", "TICKET", "CONTRACTNO", "合同号票名"],
    "品名英文":   ["品名英文", "英文品名", "品名", "产品英文名", "DESCRIPTIONS", "DESCRIPTION",
                   "COMMODITY", "品名英文COMMODITY"],
    "厂家英文":   ["厂家英文", "制造商英文", "厂家名称英文", "厂家", "制造商", "MANUFACTURER"],
    "件数":       ["件数", "总件数", "包数", "件数合计", "PACKAGES", "PACKAGE", "TOTAL", "PACKAGES件数"],
    "发货人":     ["发货人", "发货人英文", "SHIPPER"],
    "发货人地址": ["发货人地址", "SHIPPERADDRESS", "SHIPPERSADDRESS"],
    "收货人":     ["收货人", "收货人英文", "CONSIGNEE"],
    "收货人地址": ["收货人地址", "CNEEADDRESS", "CNEESADDRESS"],
    "原产国":     ["原产国", "产地", "原产地", "ORIGIN"],
}

# 输入表里允许出现、但唛头用不到的列（PL箱单 自带列），识别后忽略、不提示
IGNORED_ALIASES = [
    "ITEMS", "ITEM", "SPECIFICATIONS", "SPECIFICATION", "MATERIAL", "QUANTITY", "QTY",
    "UNIT", "NW", "NETWEIGHT", "NETWEIGHTKG", "GW", "GROSSWEIGHT", "GROSSWEIGHTKG",
    "CBM", "品名中英文", "中文品名", "品名中文", "体积", "序号",
]

# 需要从上一行「向下继承」的字段（空则继承；可通过配置关闭）
DEFAULT_FILL_DOWN = ["票名", "厂家英文", "发货人", "发货人地址", "收货人", "收货人地址", "原产国"]

DEFAULT_REQUIRED = ["票名", "品名英文", "件数", "厂家英文",
                    "发货人", "发货人地址", "收货人", "收货人地址", "原产国"]

# Times New Roman 字符宽度（单位 = 1/1000 em），用于估算换行与单页是否溢出。
# 未列出的非 ASCII 字符：CJK/全角按 1000，其余按 500。
_TNR_W = {
    " ": 250, "!": 333, '"': 408, "#": 500, "$": 500, "%": 833, "&": 778, "'": 180,
    "(": 333, ")": 333, "*": 500, "+": 564, ",": 250, "-": 333, ".": 250, "/": 278,
    ":": 278, ";": 278, "<": 564, "=": 564, ">": 564, "?": 444, "@": 921,
    "[": 333, "\\": 278, "]": 333, "^": 469, "_": 500, "`": 333,
    "{": 480, "|": 200, "}": 480, "~": 541,
}
_TNR_W.update({c: 500 for c in "0123456789"})
_TNR_W.update(dict(zip("ABCDEFGHIJKLMNOPQRSTUVWXYZ",
                       [722, 667, 667, 722, 611, 556, 722, 722, 333, 389, 722, 611, 889,
                        722, 722, 556, 722, 667, 556, 611, 722, 722, 944, 722, 722, 611])))
_TNR_W.update(dict(zip("abcdefghijklmnopqrstuvwxyz",
                       [444, 500, 444, 500, 444, 333, 500, 500, 278, 278, 500, 278, 778,
                        500, 500, 500, 500, 333, 389, 278, 500, 500, 722, 500, 500, 444])))

# XML 1.0 不允许出现的控制字符（\t \n \r 除外，但 \t 也一并换成空格）
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

# Word 单倍行距下 Times New Roman 的行高系数（ascent+descent+linegap ≈ 1.15）
LINE_HEIGHT_FACTOR = 1.15
# 安全系数：预估高度超过「可用高度 × 该系数」就报警
OVERFLOW_SAFETY = 0.95


# ---------------------------------------------------------------- 工具

def die(msg, code=2):
    sys.stderr.write("\n[错误] " + msg + "\n")
    sys.exit(code)


class IssueSink(list):
    """结构化收集校验问题。

    ★ 只给 GUI 用：CLI 走原来的「打印 + sys.exit」，行为一字不变。
    传了 sink 之后，校验函数**不再打印、不再退出**，而是把每条问题 append 进来，
    由调用方决定「哪些能忽略、哪些要问用户、哪些必须停下」。

    severity：error（无法生成）/ confirm（要问用户）/ warn（只提示）
    """

    def add(self, severity, text, advice="", row=None, field=None, cell=None, code=None):
        self.append({"severity": severity, "text": text, "advice": advice,
                     "row": row, "field": field, "cell": cell, "code": code})

    def of(self, severity):
        return [x for x in self if x["severity"] == severity]

    def has(self, severity):
        return any(x["severity"] == severity for x in self)


def _fmt_rows(rows):
    return "、".join(str(r) for r in rows)


def norm_header(s):
    """表头归一化：只保留字母/数字/汉字，其余（空格、引号、括号、换行、冒号）全去掉，再转大写。"""
    if s is None:
        return ""
    return "".join(ch for ch in str(s) if ch.isalnum()).upper()


def xml_escape(s):
    """转义 XML 实体，并剔除 XML 1.0 不允许的控制字符（\\n 由调用方先行处理）。"""
    s = _ILLEGAL_XML.sub("", str(s))
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_text(v):
    """把单元格值统一成干净的字符串：去首尾空白、统一换行符、剔非法控制字符。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        s = "TRUE" if v else "FALSE"
    elif isinstance(v, int):
        s = str(v)
    elif isinstance(v, float):
        s = str(int(v)) if v.is_integer() else repr(v)
    elif isinstance(v, str):
        s = v
    else:
        s = str(v)                      # rich text / 日期 / 其它对象
    s = _ILLEGAL_XML.sub("", s)
    s = s.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    return s.strip()


def make_wt(text):
    """把一行文字转成 <w:t>/<w:br/> 片段：\\n → 真正的换行符 <w:br/>。
    只在片段有前后空格时才加 xml:space="preserve"（与母本写法保持一致）。"""
    segs = str(text).split("\n")
    out = []
    for i, s in enumerate(segs):
        if i:
            out.append("<w:br/>")
        attr = ' xml:space="preserve"' if s != s.strip() else ""
        out.append("<w:t%s>%s</w:t>" % (attr, xml_escape(s)))
    return "".join(out)


def upper_ascii(s):
    """只把 ASCII 字母变大写，中文与其他字符原样。"""
    return "".join(ch.upper() if ("a" <= ch <= "z") else ch for ch in str(s))


def text_width_em(s):
    """按 Times New Roman 字宽表估算字符串宽度（单位 em）。"""
    w = 0
    for ch in str(s):
        if ch in _TNR_W:
            w += _TNR_W[ch]
        elif ord(ch) > 0x2E80:          # CJK / 全角 / 日韩
            w += 1000
        else:                            # × ’ — … 等符号
            w += 500
    return w / 1000.0


def is_blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def to_int(v):
    """整数解析：返回 (值 或 None, 错误说明 或 None)。"""
    if isinstance(v, bool):
        return None, "布尔值"
    if isinstance(v, (int, float)):
        n = int(v)
        if abs(float(v) - n) > 1e-9:
            return None, "不是整数（%r）" % v
        return n, None
    if isinstance(v, str):
        t = v.strip()
        if t.lstrip("+-").isdigit():
            return int(t), None
        return None, "不是整数（%r）" % v
    return None, "不是整数（%r）" % v


# ---------------------------------------------------------------- 读模板表

def parse_template(xlsx_path, cfg):
    import openpyxl

    header_row = int(cfg.get("header_row", 1))
    data_start = int(cfg.get("data_start_row", header_row + 1))
    sheet_name = cfg.get("sheet_name") or None
    fill_down = cfg.get("fill_down_fields", DEFAULT_FILL_DOWN)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    wbf = openpyxl.load_workbook(xlsx_path, data_only=False)   # 只为识别「公式没有缓存值」
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            die("模板表 %s 里没有工作表「%s」。现有：%s"
                % (os.path.basename(xlsx_path), sheet_name, "、".join(wb.sheetnames)))
        ws, wsf = wb[sheet_name], wbf[sheet_name]
    else:
        ws, wsf = wb[wb.sheetnames[0]], wbf[wbf.sheetnames[0]]

    # 1) 表头 -> 标准字段 -> 列号
    alias2std = {}
    for std, names in FIELD_ALIASES.items():
        for n in names:
            alias2std[norm_header(n)] = std
    ignored = {norm_header(n) for n in IGNORED_ALIASES}

    colmap, unknown, last_col = {}, [], 0
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(header_row, c).value
        if is_blank(raw):
            continue
        last_col = c
        key = norm_header(raw)
        std = alias2std.get(key)
        if std is None:
            if key not in ignored:
                unknown.append("%s(%s)" % (str(raw).strip(), ws.cell(header_row, c).coordinate))
            continue
        if std in colmap:
            die("表头里「%s」出现了两次（第 %d、%d 列）。" % (std, colmap[std], c))
        colmap[std] = c

    if not colmap:
        die("第 %d 行没读到任何可识别的表头。" % header_row)
    if unknown:
        # 宽容处理：用不到的列只提示、不报错（同事的表里常有多余的列）
        sys.stderr.write("[提示] 表头里有 %d 列程序用不到，已跳过（不影响生成）：%s\n"
                         % (len(unknown), "、".join(unknown)))

    # 2) 逐行读数据（值统一 to_text 归一化；同时识别「公式但没有缓存值」的格）
    rows, last = [], {}          # last = 向下继承的缓存
    for r in range(data_start, ws.max_row + 1):
        vals, formula = {}, {}
        for std, c in colmap.items():
            v = to_text(ws.cell(r, c).value)
            if not v:
                fr = wsf.cell(r, c).value
                if isinstance(fr, str) and fr.lstrip().startswith("="):
                    formula[std] = fr.strip()
            vals[std] = v
        if all(not v for v in vals.values()):
            continue                                   # 整行空 → 跳过
        # 向下继承
        for f in fill_down:
            if f in vals and not vals[f] and f in last:
                vals[f] = last[f]
        for f in fill_down:
            if f in vals and vals[f]:
                last[f] = vals[f]
        cells = {std: ws.cell(r, c).coordinate for std, c in colmap.items()}
        rows.append({"row": r, "values": vals, "cells": cells, "formula": formula})

    if not rows:
        die("数据区（第 %d 行起）没有任何数据。" % data_start)

    # 3) 按票名分组（保持首次出现顺序）
    groups, index = [], {}
    for it in rows:
        name = it["values"].get("票名")
        name = "" if is_blank(name) else str(name).strip()
        if name not in index:
            index[name] = {"name": name, "rows": []}
            groups.append(index[name])
        index[name]["rows"].append(it)

    # 票名不连续（中间夹着别的票）→ 只提示，仍按「同名归为一票」处理
    for g in groups:
        rs = [it["row"] for it in g["rows"]]
        if rs != list(range(rs[0], rs[0] + len(rs))):
            sys.stderr.write(
                "[提示] 票「%s」的行不连续（出现在第 %s 行）—— 已按同名归为一票，"
                "编号在该票内连续；若你本意是两票，请把票名写成两个不同的名字。\n"
                % (g["name"], "、".join(str(r) for r in rs)))

    return groups, last_col


# ---------------------------------------------------------------- 校验

def validate(groups, cfg, sink=None):
    """逐行校验必填项与件数。

    sink=None → CLI 老路径：问题打到 stderr 然后 sys.exit(2)。
    sink=IssueSink() → 只收集结构化问题，不打印不退出（GUI 用）。
    """
    required = cfg.get("required_fields", DEFAULT_REQUIRED)
    upper_fields = cfg.get("uppercase_fields", ["品名英文", "厂家英文"])
    errors = []

    def add(text, row=None, field=None, cell=None, advice="", code=None):
        errors.append({"text": text, "row": row, "field": field,
                       "cell": cell, "advice": advice, "code": code})

    for g in groups:
        no_name = not g["name"]
        if no_name:
            r0 = g["rows"][0]["row"]
            c0 = g["rows"][0]["cells"].get("票名")
            add("第 %d 行的「票名」为空%s，且上面也没有可继承的票名 —— 整票无法归属。"
                % (r0, "（单元格 %s）" % c0 if c0 else ""),
                row=r0, field="票名", cell=c0, code="no_ticket",
                advice="在第 %d 行的「票名」列%s里填上这一票的名字（例如 DEMO-20260101-1A）。"
                       "票名只写每票第一行即可，下面留空会自动继承。"
                       % (r0, "（%s）" % c0 if c0 else ""))
        total = 0
        for it in g["rows"]:
            tag = "第 %d 行" % it["row"]
            for f in required:
                if f == "票名" and no_name:
                    continue                       # 已按整票报过一次，不逐行刷屏
                if f not in it["values"]:
                    add("%s缺字段「%s」（表头里没有这一列）。" % (tag, f),
                        row=it["row"], field=f, code="no_column",
                        advice="在模板表第 %d 行（表头行）加一列表头「%s」，"
                               "然后把内容填在下面各行。" % (int(cfg.get("header_row", 1)), f))
                    continue
                if not it["values"][f]:
                    cell = it["cells"].get(f, "?")
                    if f in it["formula"]:
                        add("%s的「%s」是公式且没有缓存值（%s = %s）—— "
                            "请用 Excel/WPS 打开该表另存一次，让公式算出结果后再生成。"
                            % (tag, f, cell, it["formula"][f]),
                            row=it["row"], field=f, cell=cell, code="formula_no_cache",
                            advice="用 Excel 或 WPS 打开模板表 → 按 Ctrl+S 保存一次 "
                                   "（让公式算出结果）→ 再回到本程序重新生成。")
                    else:
                        add("%s的「%s」为空（单元格 %s）。" % (tag, f, cell),
                            row=it["row"], field=f, cell=cell, code="blank",
                            advice="在单元格 %s 里填上「%s」的内容。" % (cell, f))
            if it["values"].get("件数"):
                n, err = to_int(it["values"]["件数"])
                cell = it["cells"].get("件数", "?")
                if err:
                    add("%s的件数%s（单元格 %s）。" % (tag, err, cell),
                        row=it["row"], field="件数", cell=cell, code="bad_qty",
                        advice="把单元格 %s 改成 1 以上的整数（例如 2）。" % cell)
                elif n <= 0:
                    add("%s的件数必须 ≥ 1，实际 %d（单元格 %s）。" % (tag, n, cell),
                        row=it["row"], field="件数", cell=cell, code="bad_qty",
                        advice="把单元格 %s 改成 1 以上的整数（例如 2）。" % cell)
                else:
                    it["n_pkg"] = n
                    total += n
            # 大写（值已在读入时 to_text 归一化）
            for f in upper_fields:
                v = it["values"].get(f)
                if v:
                    it["values"][f] = upper_ascii(v)
        g["total_pkg"] = total

    if errors:
        if sink is not None:
            for e in errors:
                sink.add("error", e["text"], e["advice"], e["row"], e["field"],
                         e["cell"], e["code"])
            return groups
        sys.stderr.write("\n校验未通过，共 %d 处问题：\n" % len(errors))
        for i, e in enumerate(errors, 1):
            sys.stderr.write("  %2d) %s\n" % (i, e["text"]))
        sys.stderr.write("\n未生成任何文件。\n")
        sys.exit(2)
    return groups


# ---------------------------------------------------------------- PL 校验

def check_pl(groups, cfg, sink=None):
    """把每票的件数合计与同目录「中国三单/越南四单」PL箱单的 Packages 合计比对。

    sink=None → CLI 老路径（不一致就 sys.exit(2)）；
    sink=IssueSink() → 不一致记成 `confirm`（问用户要不要继续），缺文件记成 `warn`。
    """
    pc = cfg.get("pl_check") or {}
    if not pc.get("enabled", False):
        return
    raw = pc.get("dir") or "."
    dirs = raw if isinstance(raw, list) else [raw]
    bases = []
    for d in dirs:
        bases.append(d if os.path.isabs(d) else os.path.join(cfg["_config_dir"], d))
    patterns = pc.get("patterns") or ["中国三单{票名}.xlsx", "越南四单{票名}.xlsx"]
    sheet = pc.get("sheet", "PL箱单")
    pkg_col = int(pc.get("packages_col", 7))
    on_missing = pc.get("on_missing_file", "warn")

    import openpyxl
    problems, checked, missing = [], 0, []
    for g in groups:
        for pat in patterns:
            fn = pat.format(票名=g["name"], name=g["name"])
            # dir 给多个时取第一个存在的；都没有则按最后一个报「未找到」
            path = None
            for b in bases:
                cand = os.path.join(b, fn)
                if os.path.exists(cand):
                    path = cand
                    break
            if path is None:
                path = os.path.join(bases[-1], fn)
            if not os.path.exists(path):
                if on_missing == "error":
                    problems.append("找不到 %s（票「%s」）" % (os.path.basename(path), g["name"]))
                elif on_missing == "warn":
                    missing.append(os.path.basename(path))
                continue
            try:
                wb = openpyxl.load_workbook(path, data_only=True)
            except Exception as e:
                problems.append("无法读取 %s：%s" % (os.path.basename(path), e))
                continue
            if sheet not in wb.sheetnames:
                problems.append("%s 里没有工作表「%s」" % (os.path.basename(path), sheet))
                continue
            ws = wb[sheet]
            tot = 0
            for r in range(1, ws.max_row + 1):
                if isinstance(ws.cell(r, 1).value, int):
                    v = ws.cell(r, pkg_col).value
                    if isinstance(v, (int, float)):
                        tot += int(v)
            checked += 1
            if tot != g["total_pkg"]:
                problems.append("票「%s」件数合计 %d ≠ %s 的 Packages 合计 %d"
                                % (g["name"], g["total_pkg"], os.path.basename(path), tot))
            elif sink is None:
                print("  [PL校验] %s  Packages=%d  与件数合计一致 √"
                      % (os.path.basename(path), tot))

    if problems:
        if sink is not None:
            for p in problems:
                sink.add("confirm", p,
                         advice="这条只是「顺手帮你核对」：模板表的件数合计与三单/四单 PL箱单 的 "
                                "Packages 合计不一致。如果确实该不一样（例如 PL 还没更新），"
                                "选「继续生成」即可；否则先去改模板表或 PL。",
                         code="pl_mismatch")
        else:
            sys.stderr.write("\n件数一致性校验未通过：\n")
            for p in problems:
                sys.stderr.write("  - %s\n" % p)
            sys.stderr.write("\n未生成任何文件。\n")
            sys.exit(2)
    if checked == 0:
        if missing:
            if sink is not None:
                sink.add("warn", "没找到任何 PL 箱单（%s），跳过件数一致性校验。"
                         % "、".join(missing[:6]),
                         advice="不影响生成。想让程序帮你核对件数的话，"
                                "把模板表和三单/四单放在同一个文件夹里即可。",
                         code="pl_no_file")
            else:
                sys.stderr.write("[提示] 未找到任何 PL 箱单（%s），跳过件数一致性校验。\n"
                                 % "、".join(missing[:6]))
        else:
            sys.stderr.write("[提示] 没有任何 PL 文件参与校验。\n")


# ---------------------------------------------------------------- 单页溢出检查

def check_overflow(groups, labels, geom, cfg, singular=True, sink=None):
    """估算每个唛头块占多少行，超出一页可用高度就报错/警告。

    ★ 程序不会重排页面（没有渲染引擎），只能「估」：
      按母本的字号 + 页边距 + Times New Roman 字宽表算每个块占多少视觉行，
      和单页容量比。估超了就给两种处理：
        error  → 停下不产出；
        warn   → 照常产出，但把「哪几个唛头、超几行、要删几个字符」全列出来（默认）。
    估的是**整块会不会挤到第二页**；真挤过去的话，那个唛头会被拆到两页，
    而且因为它后面跟着硬分页符，后面的唛头会整体错位。
    sink=IssueSink() 时只收集，不打印不退出（GUI 用）。
    """
    oc = cfg.get("overflow_check") or {}
    if not oc.get("enabled", True):
        return
    mode = oc.get("on_overflow", "warn")
    worst, bad = 0, []
    for g in groups:
        idx = 1
        for it in g["rows"]:
            for _ in range(it["n_pkg"]):
                n, per_line = estimate_block_detail(labels, it["values"], g["total_pkg"],
                                                    idx, geom, singular)
                worst = max(worst, n)
                if n > geom["max_lines"]:
                    # 找出最「占地方」的那一段，好告诉用户去改哪里
                    lbl = max(per_line, key=lambda t: t[2])
                    bad.append((g["name"], g["total_pkg"], idx, n, it["row"],
                                lbl[1], n - geom["max_lines"]))
                idx += 1
    print("  单页容量：可用 %.0fpt / 行高 %.1fpt → 约 %.1f 行；本批最大预估 %.1f 行（余量 %.0f%%）"
          % (geom["usable_h_pt"], geom["line_h_pt"], geom["max_lines"], worst,
             (1 - worst / geom["max_lines"]) * 100))
    if not bad:
        return

    msgs, advices = [], []
    for a, b, c, d, e, lbl, over in bad:
        msgs.append("票「%s」的唛头 %d-%d（源第 %d 行）预估 %.1f 行 > 单页容量 %.1f 行，"
                    "会溢出到第二页" % (a, b, c, e, d, geom["max_lines"]))
        cut = chars_to_cut(over, geom)
        advices.append("票「%s」的唛头 %d-%d：最长的是「%s」这一段（多占 %.1f 行）—— "
                       "把它缩短约 %d 个字符（英文按字母算），或者删掉地址里的多余空格/换行，"
                       "这个唛头就能收回一页。"
                       % (a, b, c, lbl, over, cut))
    advice = ("处理办法：① 按下面点名的段落去精简（最快）；"
              "② 把配置 overflow_check.on_overflow 改成 \"warn\" 强行产出；"
              "③ 改母本页边距 / 字号后重出。")
    if mode == "error":
        if sink is not None:
            for m, ad in zip(msgs, advices):
                sink.add("error", m, ad + "（当前配置是 error 模式，不产出）", code="overflow")
            return
        sys.stderr.write("\n单页溢出检查未通过（共 %d 个唛头）：\n" % len(bad))
        for m in msgs:
            sys.stderr.write("  - %s\n" % m)
        sys.stderr.write("\n未生成任何文件。\n")
        for ad in advices:
            sys.stderr.write("  · %s\n" % ad)
        sys.stderr.write("  也可以：把配置 overflow_check.on_overflow 改成 \"warn\" 强行产出，"
                         "或改母本页边距 / 字号后重出。\n")
        sys.exit(2)
    if sink is not None:
        for m, ad in zip(msgs, advices):
            sink.add("warn", m, ad, code="overflow")
        return
    sys.stderr.write("\n[警告] 单页溢出风险（共 %d 个唛头，仅提示不阻断）：\n" % len(bad))
    for m in msgs:
        sys.stderr.write("  - %s\n" % m)
    for ad in advices:
        sys.stderr.write("  · %s\n" % ad)
    sys.stderr.write("  也可以：把 on_overflow 改成 \"error\" 让它直接停下，"
                     "或改母本页边距 / 字号。\n")


# ---------------------------------------------------------------- 生成 docx

def load_master(path):
    try:
        z = zipfile.ZipFile(path)
        d = z.read("word/document.xml").decode("utf-8")
    except zipfile.BadZipFile:
        die("母本 %s 不是有效的 docx（zip）文件。" % path)
    except KeyError:
        die("母本 %s 里没有 word/document.xml —— 它不是 Word 文档，"
            "或者是个 .xlsx。请把 master_docx 指向唛头母本 docx。" % path)
    except Exception as e:
        die("无法读取母本 %s：%s" % (path, e))

    paras = re.findall(r"<w:p[ >].*?</w:p>", d, re.S)
    if len(paras) < BLOCK_LINES + 2:
        die("母本 %s 的段落数不足（%d），不像标准唛头母本。" % (path, len(paras)))

    # 块模板校验 + 取各段「标签」（= 首个「: 」及其之前的部分）
    # ★ 只校验文字与标签，**不校验 run 个数** —— Word/WPS 重新保存会把 run 拆散，
    #   渲染是按文字位置定位的，run 怎么拆都不影响。
    block = paras[:BLOCK_LINES]
    labels = []
    for i in range(BLOCK_LINES):
        txt = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", block[i], re.S))
        exp = EXPECTED_LABELS[i]
        if exp is None:
            if txt.strip():
                die("母本第 1 块第 %d 段应为空行，实际 %r —— 母本结构不符。" % (i + 1, txt[:40]))
            labels.append("")
            continue
        if exp not in txt:
            die("母本第 1 块第 %d 段不是以「%s」开头（实际 %r）—— 母本结构不符。"
                % (i + 1, exp, txt[:40]))
        j = txt.find(": ")
        if j < 0:
            die("母本第 1 块第 %d 段找不到「: 」标签分隔（实际 %r）—— 母本结构不符。"
                % (i + 1, txt[:40]))
        if not re.search(r"<w:r(?: [^>]*)?>.*?</w:r>", block[i], re.S):
            die("母本第 1 块第 %d 段没有任何 run —— 母本结构不符。" % (i + 1))
        labels.append(txt[:j + 2])

    sep = None
    for p in paras[BLOCK_LINES:]:
        if 'w:type="page"' in p:
            sep = p
            break
    if sep is None:
        die("母本里找不到含分页符的段落，无法在块间分页。")
    tail = paras[-1]
    if 'w:type="page"' in tail:
        die("母本最后一个段落含分页符，不符合「末块无分页符」的约定。")

    body = d[d.index("<w:body>") + len("<w:body>"):]
    m = re.search(r"<w:sectPr[^>]*>.*?</w:sectPr>", body, re.S)
    if not m:
        die("母本里找不到 <w:sectPr>。")
    sectpr = m.group(0)

    # 页面几何 + 字号（用于单页溢出估算）
    def _attr(tag, name, default):
        mm = re.search(r'<%s[^>]*\bw:%s="(-?\d+)"' % (tag, name), sectpr)
        return int(mm.group(1)) if mm else default

    page_w, page_h = _attr("w:pgSz", "w", 16838), _attr("w:pgSz", "h", 11906)
    mt, mr = _attr("w:pgMar", "top", 1985), _attr("w:pgMar", "right", 720)
    mb, ml = _attr("w:pgMar", "bottom", 720), _attr("w:pgMar", "left", 720)
    szm = re.search(r'<w:sz w:val="(\d+)"', block[0])
    geom = {
        "usable_w_pt": (page_w - ml - mr) / 20.0,
        "usable_h_pt": (page_h - mt - mb) / 20.0,
        "font_pt": (int(szm.group(1)) / 2.0) if szm else 18.0,
    }
    geom["line_h_pt"] = geom["font_pt"] * LINE_HEIGHT_FACTOR
    geom["max_lines"] = geom["usable_h_pt"] * OVERFLOW_SAFETY / geom["line_h_pt"]

    return z, d[:d.index("<w:body>") + len("<w:body>")], block, sep, tail, sectpr, labels, geom


def block_texts(values, total_pkg, singular=True):
    """一个唛头块 11 行的「值部分」（行序 -> 文字）。"""
    word = "PACKAGE" if (total_pkg == 1 and singular) else "PACKAGES"
    return {
        0: values["发货人"],
        1: values["发货人地址"],
        3: values["收货人"],
        4: values["收货人地址"],
        6: values["品名英文"],
        7: values["厂家英文"],
        8: values["原产国"],
        9: "%d %s" % (total_pkg, word),
    }


def estimate_block_lines(labels, values, total_pkg, idx, geom, singular=True):
    """估算一个唛头块占多少「视觉行」（含因过长而换行、以及值里的 \\n 硬换行）。"""
    total, _ = estimate_block_detail(labels, values, total_pkg, idx, geom, singular)
    return total


def estimate_block_detail(labels, values, total_pkg, idx, geom, singular=True):
    """返回 (总行数, 每段的 (段序, 标签, 该段占的行数))，用来告诉用户「是哪一段太长」。"""
    texts = block_texts(values, total_pkg, singular)
    texts[10] = (" NO: %d" % total_pkg) + "-" + str(idx)
    total, per_line = 0, []
    for i in range(BLOCK_LINES):
        txt = labels[i] + texts.get(i, "")
        if not txt:
            total += 1                       # 空行也占一行
            per_line.append((i, labels[i].strip(": "), 1))
            continue
        n = 0
        for seg in txt.split("\n"):
            n += max(1, math.ceil(
                text_width_em(seg) * geom["font_pt"] / geom["usable_w_pt"]))
        total += n
        per_line.append((i, labels[i].strip(": "), n))
    return total, per_line


def chars_to_cut(overflow_lines, geom):
    """要把超出的 K 行收回来，大约需要删掉多少个字符（按 Times New Roman 平均字宽估）。"""
    if overflow_lines <= 0:
        return 0
    # 平均字符宽 ≈ 0.55em（大写英文字母为主的地址串）
    avg_pt = 0.55 * geom["font_pt"]
    return int(math.ceil(overflow_lines * geom["usable_w_pt"] / avg_pt))


RUN_RE = re.compile(r"<w:r(?: [^>]*)?>.*?</w:r>", re.S)


def _run_text(run):
    return "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", run, re.S))


def _run_plain(run):
    """run 的内容除 rPr / w:t 外没有别的（这种 run 可以安全重写或删除）。"""
    inner = re.sub(r"^<w:r(?: [^>]*)?>", "", run)
    inner = re.sub(r"</w:r>$", "", inner)
    inner = re.sub(r"<w:rPr>.*?</w:rPr>", "", inner, count=1, flags=re.S)
    inner = re.sub(r"<w:t[^>]*>.*?</w:t>", "", inner, flags=re.S)
    return inner.strip() == ""


def _rebuild_run(run, text):
    head = re.match(r"<w:r(?: [^>]*)?>", run).group(0)
    m = re.search(r"<w:rPr>.*?</w:rPr>", run, re.S)
    return "%s%s%s</w:r>" % (head, m.group(0) if m else "", make_wt(text))


def render_para(par, label, value):
    """把段落里「标签之后」的部分换成 value。

    ★ 不按固定的 run 下标替换 —— 而是按**文字位置**定位标签末尾，
      这样 Word/WPS 重新保存导致 run 被拆散时也不会出错。
      标签前的 run 原样保留；承载值的那一个 run 重写；其后多余的值 run 删掉。
    """
    if not label:
        return par
    runs = list(RUN_RE.finditer(par))
    if not runs:
        return par
    L = len(label)
    out, last, pos, placed = [], 0, 0, False
    for mt in runs:
        run = mt.group(0)
        t = _run_text(run)
        start, end = pos, pos + len(t)
        pos = end
        if end <= L:                                   # 整段属于标签 → 原样
            out.append(par[last:mt.end()])
            last = mt.end()
            continue
        if not placed:                                 # 承载值的那一个 run → 重写
            out.append(par[last:mt.start()])
            tail = t[:L - start] if start < L else ""  # 标签在本 run 内收尾的情况
            out.append(_rebuild_run(run, tail + value))
            last = mt.end()
            placed = True
        elif _run_plain(run):                          # 其后多余的纯值 run → 删掉
            out.append(par[last:mt.start()])
            last = mt.end()
        else:
            out.append(par[last:mt.end()])
            last = mt.end()
    out.append(par[last:])
    res = "".join(out)
    if not placed:                                     # 没有值 run → 追加一个
        m = re.search(r"<w:rPr>.*?</w:rPr>", runs[-1].group(0), re.S)
        res = res.replace("</w:p>", "<w:r>%s%s</w:r></w:p>"
                          % (m.group(0) if m else "", make_wt(value)))
    return res


def render_block(block, labels, values, total_pkg, idx, singular=True):
    """把块模板的 11 段按值渲染出来。"""
    texts = block_texts(values, total_pkg, singular)
    texts[10] = "%d-%d" % (total_pkg, idx)             # PACKAGE NO: <件数>-<序号>
    out = []
    for i in range(BLOCK_LINES):
        out.append(block[i] if not labels[i]
                   else render_para(block[i], labels[i], texts.get(i, "")))
    return "".join(out)


def renumber_ids(xml):
    c = {"p": [0x1A000000], "t": [0x1B000000]}

    def sp(m):
        c["p"][0] += 1
        return 'w14:paraId="%08X"' % c["p"][0]

    def st(m):
        c["t"][0] += 1
        return 'w14:textId="%08X"' % c["t"][0]

    xml = re.sub(r'w14:paraId="[0-9A-Fa-f]+"', sp, xml)
    return re.sub(r'w14:textId="[0-9A-Fa-f]+"', st, xml)


def build_docx(master_zip, head, block, sep, tail, sectpr, labels, groups, out_path, singular=True):
    chunks, total_blocks = [], 0
    for g in groups:
        idx = 1                                   # ★ 编号按行从上到下、每票从 1 起
        for it in g["rows"]:
            for _ in range(it["n_pkg"]):
                if total_blocks:
                    chunks.append(sep)
                chunks.append(render_block(block, labels, it["values"], g["total_pkg"], idx, singular))
                idx += 1
                total_blocks += 1
    chunks.append(tail)
    doc = renumber_ids(head + "".join(chunks) + sectpr + "</w:body></w:document>")

    app = master_zip.read("docProps/app.xml").decode("utf-8")
    app = re.sub(r"<Pages>\d+</Pages>", "<Pages>%d</Pages>" % total_blocks, app)
    app = re.sub(r"<Paragraphs>\d+</Paragraphs>",
                 "<Paragraphs>%d</Paragraphs>" % (total_blocks * 12 + 1), app)

    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in master_zip.infolist():
            name = item.filename
            if name == "word/document.xml":
                zo.writestr(item, doc.encode("utf-8"))
            elif name == "docProps/app.xml":
                zo.writestr(item, app.encode("utf-8"))
            else:
                zo.writestr(item, master_zip.read(name))
    return total_blocks


# ---------------------------------------------------------------- 主流程

def main():
    here = base_dir()
    ap = argparse.ArgumentParser(
        prog=APP_NAME, description="%s v%s —— 从 PL箱单 列一键生成 Word 唛头" % (APP_NAME, APP_VERSION))
    ap.add_argument("-V", "--version", action="version",
                    version="%s %s" % (APP_NAME, APP_VERSION))
    ap.add_argument("-c", "--config", default=os.path.join(here, "mark_config.json"))
    ap.add_argument("-i", "--input", help="模板表 xlsx（覆盖配置）")
    ap.add_argument("-o", "--out", help="输出 docx（覆盖配置）")
    ap.add_argument("--no-pl-check", action="store_true", help="跳过件数 vs PL 校验")
    ap.add_argument("--dry-run", action="store_true", help="只解析与校验，不写文件")
    ap.add_argument("--self-test", action="store_true",
                    help="跑一遍「测试用例」目录里的全部用例并校验产出")
    args = ap.parse_args()

    if args.self_test:
        return run_self_test(here)

    if not os.path.exists(args.config):
        fallback = find_asset(os.path.basename(args.config))
        if fallback and os.path.abspath(fallback) != os.path.abspath(args.config):
            sys.stderr.write("[提示] 旁边没有 %s，改用打包进程序的那份默认配置。\n"
                             % os.path.basename(args.config))
            args.config = fallback
        else:
            die("配置文件不存在：%s" % args.config)
    with open(args.config, encoding="utf-8") as f:
        cfg = json.load(f)
    # ★ 相对路径一律相对「exe 所在目录」解析（用户旁边那份配置就地生效）；
    #   只有当配置是从 exe 内部兜底取出来的，才用 exe 目录兜底。
    cfg["_config_dir"] = base_dir() if getattr(sys, "frozen", False) \
        else os.path.dirname(os.path.abspath(args.config))

    def resolve(p, default=None):
        if not p:
            return default
        return p if os.path.isabs(p) else os.path.join(cfg["_config_dir"], p)

    in_xlsx = resolve(args.input or cfg.get("input_xlsx"))
    master = resolve(cfg.get("master_docx", "唛头母本.docx"))
    if not os.path.exists(master):
        master = find_asset(os.path.basename(cfg.get("master_docx", "唛头母本.docx"))) or master
    if not in_xlsx or not os.path.exists(in_xlsx):
        die("模板表不存在：%s" % in_xlsx)
    if not os.path.exists(master):
        die("母本 docx 不存在：%s\n"
            "  → 「%s」必须和本程序放在同一个文件夹里，别只把 exe 单独拷走。"
            % (master, os.path.basename(cfg.get("master_docx", "唛头母本.docx"))))

    print("模板表：%s" % in_xlsx)
    print("母本  ：%s" % master)

    groups, _ = parse_template(in_xlsx, cfg)
    validate(groups, cfg)

    print("\n解析到 %d 个票 / %d 行明细：" % (len(groups), sum(len(g["rows"]) for g in groups)))
    total = 0
    for g in groups:
        print("  ● %s  件数合计 %d" % (g["name"], g["total_pkg"]))
        idx = 1
        for it in g["rows"]:
            v = it["values"]
            print("      第%-3d行  件数 %-3d  → NO %d-%d … %d-%d   %s / %s"
                  % (it["row"], it["n_pkg"], g["total_pkg"], idx,
                     g["total_pkg"], idx + it["n_pkg"] - 1,
                     v["品名英文"], v["厂家英文"]))
            idx += it["n_pkg"]
        total += g["total_pkg"]

    if not args.no_pl_check:
        print("\n件数一致性校验：")
        check_pl(groups, cfg)

    singular = bool(cfg.get("total_singular", True))
    master_zip, head, block, sep, tail, sectpr, labels, geom = load_master(master)

    print("\n单页溢出检查：")
    check_overflow(groups, labels, geom, cfg, singular)

    if args.dry_run:
        print("\n[dry-run] 将生成 %d 个唛头（%d 页），未写文件。" % (total, total))
        return 0

    out = args.out or cfg.get("output_docx")
    if out:
        out = resolve(out)
    else:
        # 默认命名：唛头 + 第一个票名去掉末尾分票字母（DEMO-20260101-1A → 唛头DEMO-20260101-1.docx）
        first = groups[0]["name"]
        stem = re.sub(r"(?<=[0-9])[A-Za-z]+$", "", first) or first
        out = os.path.join(cfg["_config_dir"], "唛头%s.docx" % stem)

    n = build_docx(master_zip, head, block, sep, tail, sectpr, labels, groups, out, singular=singular)
    print("\n已生成：%s" % out)
    print("  %d 个唛头 / %d 页（一件一个、块间分页符）" % (n, n))
    return 0


# ---------------------------------------------------------------- 自检（--self-test）

def parse_mark_docx(path):
    """把一个唛头 docx 拆成块，用于自检。"""
    import xml.etree.ElementTree as ET
    d = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    try:
        ET.fromstring(d)
        xml_ok = True
    except ET.ParseError:
        xml_ok = False
    paras = re.findall(r"<w:p[ >].*?</w:p>", d, re.S)
    txt = ["".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", pa, re.S)) for pa in paras]
    blocks, cur = [], []
    for pa, t in zip(paras, txt):
        if 'w:type="page"' in pa:
            blocks.append(cur)
            cur = []
            continue
        cur.append(t)
    blocks.append(cur)
    while blocks and blocks[-1] and blocks[-1][-1] == "":
        blocks[-1].pop()
    out = []
    for b in blocks:
        if not b or not b[0].startswith("SHIPPER:"):
            continue
        g = lambda k: next((x.split(":", 1)[1].strip() for x in b if x.startswith(k + ":")), "")
        out.append({"commodity": g("COMMODITY"), "maker": g("MANUFACTURER"),
                    "total": g("TOTAL"), "no": g("PACKAGE NO"),
                    "lines": len([x for x in b if x != ""]) + 2})
    ids = re.findall(r'w14:paraId="([0-9A-Fa-f]+)"', d)
    return out, d.count('w:type="page"'), len(ids) == len(set(ids)), xml_ok, d


def run_self_test(here):
    """跑「测试用例」目录里的全部用例（`负面*` = 期望失败）并逐项断言。"""
    cases_dir = os.path.join(here, "测试用例")
    files = sorted(f for f in glob.glob(os.path.join(cases_dir, "*.xlsx"))
                   if not os.path.basename(f).startswith("~$"))
    if not files:
        sys.stderr.write("[错误] 找不到测试用例：%s\n" % cases_dir)
        return 2
    print("%s %s —— 自检（%d 个用例）\n" % (APP_NAME, APP_VERSION, len(files)))
    # 调自己跑每个用例：打包成 exe 时 sys.executable 就是 exe；
    # 未打包时要显式带上脚本路径，否则 "-i" 会被 Python 当成「交互模式」而挂住等 stdin。
    frozen = bool(getattr(sys, "frozen", False))
    self_cmd = [sys.executable] + ([] if frozen else [os.path.join(here, "make_mark.py")])
    # 子进程 stdout 是管道（非控制台）时，Python 会退回本地代码页（中文 Windows = cp936），
    # 若父进程按 utf-8 严格解码就会在 reader 线程里抛 UnicodeDecodeError。
    # 双保险：① 让子进程强制 utf-8 输出；② 父进程容错解码。
    child_env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
    ok_all = True
    for x in files:
        name = os.path.splitext(os.path.basename(x))[0]
        negative = name.startswith("负面")
        out = os.path.join(cases_dir, "唛头_%s.docx" % name)
        if os.path.exists(out):
            os.remove(out)
        print("=" * 92)
        print("■ %s%s" % (name, "（负面用例：期望报错停下）" if negative else ""))
        r = subprocess.run(self_cmd + ["-i", x, "-o", out, "--no-pl-check"],
                           capture_output=True, text=True, encoding="utf-8",
                           errors="replace", env=child_env,
                           stdin=subprocess.DEVNULL)
        body = (r.stdout or "") + (r.stderr or "")
        if negative:
            good = (r.returncode != 0) and (not os.path.exists(out))
            print("   退出码 %d（应非 0）%s | 未产出文件 %s"
                  % (r.returncode, "√" if r.returncode != 0 else "×",
                     "√" if not os.path.exists(out) else "×"))
            for l in body.splitlines():
                if "[错误]" in l or re.match(r"\s*\d+\)", l):
                    print("   " + l.strip())
            print("   %s" % ("√ 符合预期" if good else "× 与预期不符"))
            ok_all &= good
            continue
        if r.returncode != 0:
            sys.stderr.write(r.stderr)
            print("   × 生成失败，退出码 %d" % r.returncode)
            ok_all = False
            continue
        blocks, breaks, uniq, xml_ok, xml = parse_mark_docx(out)
        print("   块数 %d / 分页符 %d / paraId 唯一 %s / XML 合法 %s"
              % (len(blocks), breaks, uniq, xml_ok))
        print("    #   TOTAL         PACKAGE NO   COMMODITY                 MANUFACTURER")
        for i, b in enumerate(blocks, 1):
            print("    %-3d %-13s %-12s %-25s %s"
                  % (i, b["total"], b["no"], b["commodity"], b["maker"]))
        checks = [
            ("块数 = 分页符+1", len(blocks) == breaks + 1),
            ("每块 11 行", all(b["lines"] == 11 for b in blocks)),
            ("paraId 唯一", uniq),
            ("document.xml 合法", xml_ok),
            ("PACKAGE NO 形如 N-i", all(re.fullmatch(r"\d+-\d+", b["no"]) for b in blocks)),
            ("TOTAL 与 NO 前半段一致",
             all(b["no"].split("-")[0] == b["total"].split()[0] for b in blocks)),
            ("& 已转义成 &amp;", "&amp;" in xml or "&" not in xml),
            ("无 XML 非法控制字符", not re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", xml)),
            ("<w:t> 里没有裸换行", not re.search(r"<w:t[^>]*>[^<]*\n", xml)),
        ]
        for label, good in checks:
            print("   %s %s" % ("√" if good else "×", label))
            ok_all &= good
    print("=" * 92)
    print("自检结果：%s" % ("全部通过 √" if ok_all else "有失败项 ×"))
    return 0 if ok_all else 1


if __name__ == "__main__":
    sys.exit(main() or 0)
